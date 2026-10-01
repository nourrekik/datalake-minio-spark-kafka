"""Ingestion automatisée : incoming/ -> RAW (missions 4 à 7).

Pour chaque fichier déposé dans incoming/<source>/ :
  1. contrôle      : extension attendue, fichier non vide, format lisible, champs obligatoires présents
  2. empreinte     : calcul du hash SHA-256 du contenu
  3. anti-doublon  : si ce hash figure déjà dans le manifest -> fichier ignoré (DUPLICATE)
  4. dépôt         : PutObject dans raw/<source>/ingest_date=AAAA-MM-JJ/<fichier>, avec métadonnées
  5. manifest      : enregistrement du hash dans raw/_ingestion/manifest.json (dans MinIO)
  6. journal       : une ligne par fichier dans logs/ (date, source, fichier, destination, taille, statut, motif)
  7. rangement     : succès/doublon -> incoming/_archive/ ; fichier invalide -> incoming/_rejected/
                     MinIO injoignable ou credentials refusés -> le fichier RESTE dans incoming (reprise au prochain lancement)

Usage : python src/ingestion/ingest.py
Code retour : 0 si tout est ingéré ou ignoré, 1 si des fichiers sont rejetés, 2 si l'ingestion est impossible.
"""
import csv
import hashlib
import io
import json
import shutil
import sys
import xml.etree.ElementTree as ET
from datetime import datetime
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from botocore.exceptions import BotoCoreError, ClientError, EndpointConnectionError  # noqa: E402
from common.config import BUCKET, PROJECT_ROOT, SETTINGS, get_s3_client  # noqa: E402

INCOMING = PROJECT_ROOT / SETTINGS["incoming_dir"]
ARCHIVE = INCOMING / "_archive"
REJECTED = INCOMING / "_rejected"
JOURNAL_DIR = PROJECT_ROOT / SETTINGS["journal_dir"]
MANIFEST_KEY = SETTINGS["manifest_key"]
SOURCES = SETTINGS["sources"]
ERREURS_ACCES = {"AccessDenied", "InvalidAccessKeyId", "SignatureDoesNotMatch", "ExpiredToken"}


class FichierInvalide(Exception):
    """Le fichier lui-même pose problème : il est rejeté (pas de nouvelle tentative)."""


class IngestionImpossible(Exception):
    """Problème d'infrastructure (MinIO, credentials) : on s'arrête, les fichiers restent pour une reprise."""


# --------------------------------------------------------------------------- journal
def journaliser(source, fichier, destination, taille, statut, motif=""):
    JOURNAL_DIR.mkdir(exist_ok=True)
    maintenant = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    entree = {"date": maintenant, "source": source, "fichier": fichier, "destination": destination,
              "taille_octets": taille, "statut": statut, "motif": motif}
    with open(JOURNAL_DIR / "ingestion_journal.jsonl", "a", encoding="utf-8") as f:
        f.write(json.dumps(entree, ensure_ascii=False) + "\n")
    ligne = f"{maintenant} | {source:<8} | {fichier} | {destination or '-'} | {taille} o | {statut}"
    if motif:
        ligne += f" | {motif}"
    with open(JOURNAL_DIR / "ingestion.log", "a", encoding="utf-8") as f:
        f.write(ligne + "\n")
    print("   " + ligne)


# --------------------------------------------------------------------------- contrôles
def controler(source, chemin, contenu):
    conf = SOURCES[source]
    if chemin.suffix.lower() not in conf["extensions"]:
        raise FichierInvalide(f"extension {chemin.suffix or '(aucune)'} inattendue pour {source} "
                              f"(attendu : {', '.join(conf['extensions'])})")
    if len(contenu) == 0 or not contenu.strip():
        raise FichierInvalide("fichier vide")
    obligatoires = set(conf["champs_obligatoires"])
    ext = chemin.suffix.lower()
    try:
        if ext == ".csv":
            texte = contenu.decode("utf-8-sig")
            lignes = list(csv.reader(io.StringIO(texte)))
            champs = {c.strip() for c in lignes[0]}
            nb = len(lignes) - 1
        elif ext == ".json":
            data = json.loads(contenu.decode("utf-8-sig"))
            if not isinstance(data, list) or not data or not isinstance(data[0], dict):
                raise FichierInvalide("JSON inattendu : une liste d'objets est attendue")
            champs = set().union(*(d.keys() for d in data if isinstance(d, dict)))
            nb = len(data)
        elif ext == ".xml":
            racine = ET.fromstring(contenu)
            enregistrements = list(racine)
            if not enregistrements:
                raise FichierInvalide("XML sans enregistrement")
            champs = {e.tag for e in enregistrements[0]}
            nb = len(enregistrements)
        elif ext == ".xlsx":
            from openpyxl import load_workbook
            ws = load_workbook(io.BytesIO(contenu), read_only=True).active
            lignes = list(ws.iter_rows(values_only=True))
            champs = {str(c).strip() for c in lignes[0] if c is not None}
            nb = len(lignes) - 1
    except FichierInvalide:
        raise
    except Exception as e:  # contenu illisible : corrompu ou mauvais format
        raise FichierInvalide(f"fichier corrompu ou format illisible ({type(e).__name__}: {str(e)[:80]})")
    manquants = obligatoires - champs
    if manquants:
        raise FichierInvalide(f"champs obligatoires absents : {', '.join(sorted(manquants))}")
    if nb <= 0:
        raise FichierInvalide("aucune ligne de données")
    return nb


# --------------------------------------------------------------------------- MinIO
def appel_minio(fonction, *args, **kwargs):
    """Exécute un appel S3 en transformant les erreurs d'infrastructure en IngestionImpossible."""
    try:
        return fonction(*args, **kwargs)
    except EndpointConnectionError:
        raise IngestionImpossible("MinIO injoignable (serveur arrêté ou réseau indisponible)")
    except ClientError as e:
        code = e.response["Error"].get("Code")
        if code in ERREURS_ACCES:
            raise IngestionImpossible(f"accès refusé par MinIO ({code}) : credentials incorrects ou droits insuffisants")
        raise IngestionImpossible(f"erreur MinIO {code} : {e.response['Error'].get('Message')}")
    except BotoCoreError as e:
        raise IngestionImpossible(f"erreur de connexion à MinIO ({type(e).__name__})")


def charger_manifest(s3):
    try:
        obj = s3.get_object(Bucket=BUCKET, Key=MANIFEST_KEY)
        return json.loads(obj["Body"].read())
    except ClientError as e:
        if e.response["Error"].get("Code") in ("NoSuchKey", "404"):
            return {}  # première ingestion
        raise


def sauver_manifest(s3, manifest):
    corps = json.dumps(manifest, ensure_ascii=False, indent=2).encode("utf-8")
    appel_minio(s3.put_object, Bucket=BUCKET, Key=MANIFEST_KEY, Body=corps, ContentType="application/json")


def ranger(chemin, dossier, source):
    cible = dossier / source
    cible.mkdir(parents=True, exist_ok=True)
    destination = cible / chemin.name
    if destination.exists():  # ne jamais écraser une archive précédente
        destination = cible / f"{chemin.stem}_{datetime.now():%Y%m%d%H%M%S}{chemin.suffix}"
    shutil.move(str(chemin), destination)


# --------------------------------------------------------------------------- programme
def fichiers_a_traiter():
    for source in SOURCES:
        dossier = INCOMING / source
        dossier.mkdir(parents=True, exist_ok=True)
        for chemin in sorted(dossier.iterdir()):
            if chemin.is_file() and not chemin.name.startswith("."):
                yield source, chemin


def main():
    print(f"=== Ingestion incoming/ -> s3://{BUCKET}/raw/  ({datetime.now():%Y-%m-%d %H:%M:%S}) ===")
    a_traiter = list(fichiers_a_traiter())
    if not a_traiter:
        print("Aucun nouveau fichier dans incoming/.")
        return 0
    print(f"{len(a_traiter)} fichier(s) détecté(s).")

    bilan = {"SUCCESS": 0, "DUPLICATE": 0, "REJECTED": 0, "FAILED": 0}
    try:
        s3 = get_s3_client("ingestion")
        manifest = appel_minio(charger_manifest, s3)   # vérifie aussi connexion + credentials
    except (IngestionImpossible, FileNotFoundError) as e:
        for source, chemin in a_traiter:
            journaliser(source, chemin.name, "", chemin.stat().st_size, "FAILED", f"{e} — fichier conservé pour reprise")
        print(f"\n!! Ingestion impossible : {e}")
        return 2

    for i, (source, chemin) in enumerate(a_traiter):
        contenu = chemin.read_bytes()
        taille = len(contenu)
        empreinte = hashlib.sha256(contenu).hexdigest()
        try:
            nb = controler(source, chemin, contenu)

            if empreinte in manifest:
                deja = manifest[empreinte]
                journaliser(source, chemin.name, deja["destination"], taille, "DUPLICATE",
                            f"contenu identique déjà ingéré le {deja['date']} ({deja['fichier']})")
                bilan["DUPLICATE"] += 1
                ranger(chemin, ARCHIVE, source)
                continue

            date = datetime.now()
            cle = f"{SOURCES[source]['raw_prefix']}ingest_date={date:%Y-%m-%d}/{chemin.name}"
            existe = appel_minio(s3.list_objects_v2, Bucket=BUCKET, Prefix=cle, MaxKeys=1).get("KeyCount", 0)
            if existe:  # même nom mais contenu différent le même jour : on ne l'écrase pas
                cle = cle[: -len(chemin.suffix)] + f"__{empreinte[:8]}{chemin.suffix}"

            appel_minio(s3.put_object, Bucket=BUCKET, Key=cle, Body=contenu,
                        Metadata={"sha256": empreinte, "source": source, "fichier-origine": chemin.name,
                                  "ingere-le": date.isoformat(timespec="seconds")})
            manifest[empreinte] = {"fichier": chemin.name, "source": source, "destination": cle,
                                   "taille": taille, "lignes": nb, "date": date.strftime("%Y-%m-%d %H:%M:%S")}
            sauver_manifest(s3, manifest)
            journaliser(source, chemin.name, cle, taille, "SUCCESS", f"{nb} enregistrement(s)")
            bilan["SUCCESS"] += 1
            ranger(chemin, ARCHIVE, source)

        except FichierInvalide as e:
            journaliser(source, chemin.name, "", taille, "REJECTED", str(e))
            bilan["REJECTED"] += 1
            ranger(chemin, REJECTED, source)
        except IngestionImpossible as e:
            for src, ch in a_traiter[i:]:
                journaliser(src, ch.name, "", ch.stat().st_size, "FAILED", f"{e} — fichier conservé pour reprise")
                bilan["FAILED"] += 1
            print(f"\n!! Ingestion interrompue : {e}")
            break

    print(f"\nBilan : {bilan['SUCCESS']} ingéré(s), {bilan['DUPLICATE']} doublon(s) ignoré(s), "
          f"{bilan['REJECTED']} rejeté(s), {bilan['FAILED']} en échec (à reprendre).")
    print(f"Journal : {JOURNAL_DIR / 'ingestion.log'}")
    if bilan["FAILED"]:
        return 2
    return 1 if bilan["REJECTED"] else 0


if __name__ == "__main__":
    sys.exit(main())
